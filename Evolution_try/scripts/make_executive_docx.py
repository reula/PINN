"""Render EXECUTIVE_REPORT.md as a Word document.

Line-based parsing is not enough: the Markdown is wrapped at ~90 columns, so a bold
span routinely opens on one line and closes on the next, and a paragraph spans several
lines.  Both have to be joined BEFORE the inline markers are resolved, or the markers
survive into the document as literal asterisks.
"""
import re
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.table import WD_TABLE_ALIGNMENT

SRC, OUT = "EXECUTIVE_REPORT.md", "EXECUTIVE_REPORT.docx"
lines = open(SRC).read().split("\n")

doc = Document()
n = doc.styles["Normal"]; n.font.name = "Calibri"; n.font.size = Pt(10.5)
for s in doc.sections:
    s.left_margin = s.right_margin = Inches(0.8)
    s.top_margin = s.bottom_margin = Inches(0.8)

TOKEN = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*|\[[^\]]+\]\([^)]+\))")

def inline(par, text, bold=False, italic=False):
    """Resolve markup RECURSIVELY: the source nests code inside bold (`**rel `L2`**`),
    and a single-level pass would leave the inner backticks as literal characters."""
    for tok in TOKEN.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            inline(par, tok[2:-2], bold=True, italic=italic)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            inline(par, tok[1:-1], bold=bold, italic=True)
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            r = par.add_run(tok[1:-1]); r.font.name = "Consolas"; r.font.size = Pt(9.5)
            if bold: r.bold = True
            if italic: r.italic = True
        elif tok.startswith("["):
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", tok)
            r = par.add_run(m.group(1)); r.underline = True
            if bold: r.bold = True
        else:
            r = par.add_run(tok)
            if bold: r.bold = True
            if italic: r.italic = True

def add_table(rows):
    head, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(head)); t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(head):
        cell = t.rows[0].cells[i]; cell.text = ""
        inline(cell.paragraphs[0], c, bold=True)
        for r in cell.paragraphs[0].runs: r.font.size = Pt(8.5)
    for row in body:
        cells = t.add_row().cells
        for i, c in enumerate(row[:len(head)]):
            cells[i].text = ""
            inline(cells[i].paragraphs[0], c)
            for r in cells[i].paragraphs[0].runs: r.font.size = Pt(8.5)
    doc.add_paragraph()

buf = []          # a logical paragraph, possibly several wrapped source lines
def flush(style=None, bold=False, italic=False):
    global buf
    if buf:
        text = " ".join(x.strip() for x in buf).strip()
        if text:
            p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
            inline(p, text, bold=bold, italic=italic)
    buf = []

i = 0
while i < len(lines):
    ln = lines[i].rstrip()
    stripped = ln.strip()

    if not stripped:
        flush(); i += 1; continue

    # a table: header row, separator row, then body rows
    if stripped.startswith("|") and i + 1 < len(lines) and \
       set(lines[i + 1].replace("|", "").strip()) <= set("-: "):
        flush(); rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not (set("".join(cells)) <= set("-: ")):
                rows.append(cells)
            i += 1
        if rows:
            add_table(rows)
        continue

    if stripped.startswith("```"):
        flush(); i += 1; code = []
        while i < len(lines) and not lines[i].strip().startswith("```"):
            code.append(lines[i]); i += 1
        p = doc.add_paragraph(); r = p.add_run("\n".join(code))
        r.font.name = "Consolas"; r.font.size = Pt(8.5)
        p.paragraph_format.left_indent = Inches(0.2)
        i += 1; continue

    if stripped == "---":
        flush(); doc.add_paragraph(); i += 1; continue

    for pre, lvl in (("### ", 2), ("## ", 1), ("# ", 0)):
        if stripped.startswith(pre):
            flush(); inline(doc.add_heading("", level=lvl), stripped[len(pre):])
            i += 1; break
    else:
        if stripped.startswith("> "):
            buf.append(stripped[2:]); i += 1; continue
        if stripped.startswith("- "):
            flush(); buf = [stripped[2:]]
            flush(style="List Bullet"); i += 1; continue
        m = re.match(r"^(\d+)\. (.*)$", stripped)
        if m:
            flush(); buf = [m.group(2)]
            flush(style="List Number"); i += 1; continue
        buf.append(stripped); i += 1
        continue
    if stripped.startswith("### ") or stripped.startswith("## ") or stripped.startswith("# "):
        continue

flush()
doc.save(OUT)
print("wrote", OUT, "| paragraphs:", len(doc.paragraphs), "| tables:", len(doc.tables))
